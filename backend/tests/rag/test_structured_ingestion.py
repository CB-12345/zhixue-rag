from types import SimpleNamespace

import fitz
import pytest
from docx import Document as WordDocument
from langchain_core.documents import Document

from app.rag.document_handler.processor import DocumentProcessor
from app.rag.text_spliter import AsyncTextSplitter
from app.utils import pdf_multimodal_loader as pdf_loader
from app.utils.pdf_text import extract_page_text


@pytest.mark.parametrize("asynchronous", [False, True])
async def test_real_docx_late_facts_keep_their_heading_and_do_not_take_the_next_heading(tmp_path, asynchronous):
    path = tmp_path / "course.docx"
    word = WordDocument()
    word.add_heading("课程规则 LAB-461", 1)
    word.add_heading("设备故障规则", 2)
    for _ in range(9):
        word.add_paragraph("这是背景段落，记录观察事实并核对设备编号。" * 12)
    word.add_paragraph("负责人是江予川，发现电源损坏须暂停使用。")
    word.add_heading("另一项独立政策", 2)
    word.add_paragraph("负责人是唐亦安。")
    word.save(path)
    processor = DocumentProcessor(object())
    docs = processor.get_file_document_sync(str(path))
    chunks = (await processor.spliter.split_documents(docs) if asynchronous else processor.split_documents_sync(docs))
    target = next(c for c in chunks if "江予川" in c.page_content)
    other = next(c for c in chunks if "唐亦安" in c.page_content)
    assert target.page_content.startswith("# 课程规则 LAB-461\n## 设备故障规则")
    assert "另一项独立政策" not in target.page_content
    assert "设备故障规则" not in other.page_content
    assert all(len(c.page_content) <= 1000 and "page" not in c.metadata for c in chunks)


def test_table_rows_keep_headers_without_losing_or_mixing_values():
    rows = [f"| ID-{i:02} | {i+30} 小时 | 入口-{i:02} |" for i in range(30)]
    text = "# 申请规则\n\n| 申请编号 | 反馈时限 | 提交入口 |\n| --- | --- | --- |\n" + "\n".join(rows)
    chunks = AsyncTextSplitter(chunk_size=170, chunk_overlap=10).split_documents_sync([
        Document(page_content=text, metadata={"page": 3, "source": "a.pdf"})])
    assert len(chunks) > 5
    for row in rows:
        assert sum(row in c.page_content for c in chunks) == 1
    assert all("| 申请编号 | 反馈时限 | 提交入口 |" in c.page_content for c in chunks)
    assert all(c.page_content.startswith("# 申请规则") and len(c.page_content) <= 170 for c in chunks)
    assert all(c.metadata["page"] == 3 for c in chunks)


def test_oversized_cell_repeats_identity_and_column_label_and_preserves_value():
    long_value = "ABCDEFGHIJ" * 25
    text = "| 编号 | 说明 |\n| --- | --- |\n| ROW-1 | " + long_value + " |"
    chunks = AsyncTextSplitter(chunk_size=75, chunk_overlap=0).split_documents_sync([Document(page_content=text)])
    assert len(chunks) > 1 and all(len(c.page_content) <= 75 for c in chunks)
    prefix = "编号：ROW-1\n说明："
    assert all(c.page_content.startswith(prefix) for c in chunks)
    assert "".join(c.page_content[len(prefix):] for c in chunks) == long_value


def test_code_fences_do_not_create_fake_sections_or_tables():
    text = "# Real section\n\n```markdown\n# Fake section\n| x | y |\n| --- | --- |\n| A | B |\n```\n\nActual text."
    chunks = AsyncTextSplitter(chunk_size=200, chunk_overlap=0).split_documents_sync([Document(page_content=text)])
    assert len(chunks) == 1
    assert chunks[0].metadata["section_path"] == ["Real section"]
    assert "# Fake section" in chunks[0].page_content
    assert chunks[0].metadata["structure_type"] == "text"


def test_headings_do_not_leak_between_documents_or_pdf_pages():
    docs = [Document(page_content="# First section\nfirst fact", metadata={"page": 1}),
            Document(page_content="second fact", metadata={"page": 2})]
    chunks = AsyncTextSplitter(chunk_size=100, chunk_overlap=5).split_documents_sync(docs)
    assert chunks[1].page_content == "second fact" and chunks[1].metadata == {"page": 2}


def test_heading_punctuation_is_source_content_not_a_closing_heading_marker():
    chunks = AsyncTextSplitter(chunk_size=100, chunk_overlap=5).split_documents_sync([
        Document(page_content="# C#\nLanguage topic.\n## Other ###\nNext topic.")])
    assert chunks[0].metadata["section_path"] == ["C#"]
    assert chunks[1].metadata["section_path"] == ["C#", "Other"]


def test_header_only_table_and_title_only_section_survive():
    text = "# 空表\n\n| A | B |\n| --- | --- |\n\n# 只有标题"
    chunks = AsyncTextSplitter(chunk_size=100, chunk_overlap=5).split_documents_sync([Document(page_content=text)])
    assert any("| A | B |\n| --- | --- |" in c.page_content for c in chunks)
    assert any(c.page_content == "# 只有标题" for c in chunks)


def test_escaped_pipes_do_not_create_false_columns():
    text = "| 编号 | 说明 |\n| --- | --- |\n| KEY | value \\| other |"
    chunks = AsyncTextSplitter(chunk_size=100, chunk_overlap=5).split_documents_sync([Document(page_content=text)])
    assert len(chunks) == 1 and "value \\| other" in chunks[0].page_content
    assert chunks[0].metadata["structure_type"] == "table"


def test_huge_title_or_table_header_does_not_drop_body_or_exceed_budget():
    for text in ["# " + "H" * 180 + "\nTAIL-BODY", "| " + "HEADER" * 20 + " | B |\n| --- | --- |\n| VALUE-A | VALUE-B |"]:
        chunks = AsyncTextSplitter(chunk_size=70, chunk_overlap=0).split_documents_sync([Document(page_content=text)])
        joined = "".join(c.page_content for c in chunks)
        assert all(len(c.page_content) <= 70 for c in chunks)
        assert "TAIL-BODY" in joined if "TAIL-BODY" in text else "VALUE-A" in joined and "VALUE-B" in joined


def make_grid_pdf(path):
    with fitz.open() as pdf:
        pdf.new_page().insert_text((45, 45), "Introduction with no table.")
        page = pdf.new_page()
        page.insert_text((45, 45), "Equipment access table. Do not confuse nearby identifiers.")
        xs, top, height = [45, 155, 290, 545], 85, 25
        rows = [["Equipment", "Open hours", "Requirement"],
                *[[f"EQ-{i:03}", f"{i+8:02}:30", f"Receipt {i:03} required"] for i in range(12)]]
        for x in xs:
            page.draw_line((x, top), (x, top + height * len(rows)))
        for i in range(len(rows) + 1):
            page.draw_line((xs[0], top + height * i), (xs[-1], top + height * i))
        for i, row in enumerate(rows):
            for x, value in zip(xs, row):
                page.insert_text((x + 5, top + height * i + 16), value, fontsize=10)
        page.insert_text((45, 470), "Outside text must also survive.")
        pdf.save(path)


@pytest.mark.parametrize("asynchronous", [False, True])
async def test_real_pdf_grid_reaches_chunks_with_columns_and_physical_page(tmp_path, monkeypatch, asynchronous):
    path = tmp_path / "grid.pdf"
    make_grid_pdf(path)
    monkeypatch.setattr(pdf_loader, "extract_images_from_pdf", lambda *args: {})
    docs = (await pdf_loader.pdf_multimodal_loader(str(path), "fixture", "test") if asynchronous
            else pdf_loader.pdf_multimodal_loader_sync(str(path), "fixture", "test"))
    assert [d.metadata["page"] for d in docs] == [1, 2]
    assert "Outside text must also survive." in docs[1].page_content
    chunks = AsyncTextSplitter(chunk_size=200, chunk_overlap=10).split_documents_sync(docs)
    table = [c for c in chunks if c.metadata.get("structure_type") == "table"]
    assert len(table) > 2
    assert all("| Equipment | Open hours | Requirement |" in c.page_content and c.metadata["page"] == 2 for c in table)
    assert sum("| EQ-008 | 16:30 | Receipt 008 required |" in c.page_content for c in table) == 1


def test_failed_table_detection_preserves_original_text():
    def unavailable(**kwargs):
        raise RuntimeError("unavailable")
    page = SimpleNamespace(get_text=lambda *args: "Original paragraph and table text.", find_tables=unavailable)
    assert extract_page_text(page) == "Original paragraph and table text."


def test_merged_grid_is_not_silently_reconstructed_with_guessed_values():
    table = SimpleNamespace(extract=lambda: [["A", "B"], ["value", None]])
    page = SimpleNamespace(get_text=lambda *args: "original merged cell text", find_tables=lambda **kwargs: SimpleNamespace(tables=[table]))
    assert extract_page_text(page) == "original merged cell text"
