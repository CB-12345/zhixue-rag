"""Unit tests for app.utils.file_handler helpers and loaders.

Real files are limited to tmp_path fixtures, including valid DOCX packages.
PDF/PPTX parsers are only exercised on error/empty paths.
"""
import hashlib
import io
import os

import pytest
from docx import Document as WordDocument

from app.utils.file_handler import (
    FontBBoxStreamFilter,
    get_file_md5_hex,
    get_file_md5_hex_sync,
    listdir_allowed_type,
    markdown_loader,
    markdown_loader_sync,
    ppt_loader,
    ppt_loader_sync,
    txt_loader,
    txt_loader_sync,
    word_loader,
    word_loader_sync,
)


# ---------------------------------------------------------------------------
# md5 helpers
# ---------------------------------------------------------------------------
def test_get_file_md5_hex_sync_matches_hashlib(tmp_path):
    p = tmp_path / "file.bin"
    p.write_bytes(bytes(range(256)) * 16)
    assert get_file_md5_hex_sync(str(p)) == hashlib.md5(p.read_bytes()).hexdigest()


async def test_get_file_md5_hex_async_matches_hashlib(tmp_path):
    p = tmp_path / "file.bin"
    p.write_bytes(b"hello world")
    assert await get_file_md5_hex(str(p)) == hashlib.md5(b"hello world").hexdigest()


def test_get_file_md5_hex_sync_missing_file_returns_empty(tmp_path):
    assert get_file_md5_hex_sync(str(tmp_path / "nope.bin")) == ""


async def test_get_file_md5_hex_async_missing_file_returns_empty(tmp_path):
    assert await get_file_md5_hex(str(tmp_path / "nope.bin")) == ""


def test_get_file_md5_hex_sync_directory_returns_empty(tmp_path):
    assert get_file_md5_hex_sync(str(tmp_path)) == ""


async def test_get_file_md5_hex_async_directory_returns_empty(tmp_path):
    assert await get_file_md5_hex(str(tmp_path)) == ""


def test_get_file_md5_hex_sync_empty_file(tmp_path):
    p = tmp_path / "empty.bin"
    p.write_bytes(b"")
    assert get_file_md5_hex_sync(str(p)) == hashlib.md5(b"").hexdigest()


# ---------------------------------------------------------------------------
# listdir_allowed_type
# ---------------------------------------------------------------------------
async def test_listdir_allowed_type_filters_by_extension(tmp_path):
    (tmp_path / "a.txt").write_text("a")
    (tmp_path / "b.pdf").write_text("b")
    (tmp_path / "c.md").write_text("c")
    (tmp_path / "d.png").write_text("d")
    (tmp_path / "e.TXT").write_text("e")  # case-sensitive -> excluded
    subdir = tmp_path / "sub"
    subdir.mkdir()
    (subdir / "f.txt").write_text("f")  # nested -> excluded

    result = await listdir_allowed_type(str(tmp_path), (".txt", ".pdf"))
    assert isinstance(result, tuple)
    names = sorted(str(p).replace("\\", "/").rsplit("/", 1)[-1] for p in result)
    assert names == ["a.txt", "b.pdf"]
    # returned paths are absolute and point at real files
    for p in result:
        assert p.startswith(str(tmp_path))
        assert os.path.isfile(p)


async def test_listdir_allowed_type_nonexistent_dir_returns_empty(tmp_path):
    assert await listdir_allowed_type(str(tmp_path / "nope"), (".txt",)) == ()


async def test_listdir_allowed_type_file_path_returns_empty(tmp_path):
    p = tmp_path / "single.txt"
    p.write_text("x")
    assert await listdir_allowed_type(str(p), (".txt",)) == ()


async def test_listdir_allowed_type_empty_dir_returns_empty(tmp_path):
    assert await listdir_allowed_type(str(tmp_path), (".txt",)) == ()


# ---------------------------------------------------------------------------
# text loaders
# ---------------------------------------------------------------------------
def test_txt_loader_sync_utf8(tmp_path):
    p = tmp_path / "note.txt"
    content = "第一行\nsecond line\n第三行"
    p.write_text(content, encoding="utf-8")
    docs = txt_loader_sync(str(p))
    assert len(docs) == 1
    assert content in docs[0].page_content


def test_txt_loader_sync_gbk_fallback(tmp_path):
    p = tmp_path / "gbk.txt"
    content = "你好世界中文内容"
    p.write_text(content, encoding="gbk")
    docs = txt_loader_sync(str(p))
    assert len(docs) == 1
    assert content.strip() in docs[0].page_content


async def test_txt_loader_async_roundtrip(tmp_path):
    p = tmp_path / "note.txt"
    content = "async loader content"
    p.write_text(content, encoding="utf-8")
    docs = await txt_loader(str(p))
    assert len(docs) == 1
    assert content in docs[0].page_content


def test_txt_loader_sync_missing_file_returns_empty(tmp_path):
    assert txt_loader_sync(str(tmp_path / "nope.txt")) == []


async def test_txt_loader_async_missing_file_returns_empty(tmp_path):
    assert await txt_loader(str(tmp_path / "nope.txt")) == []


# ---------------------------------------------------------------------------
# Word: real documents, including headings, tables, empty/corrupt packages
# ---------------------------------------------------------------------------
@pytest.mark.parametrize("asynchronous", [False, True])
async def test_word_loader_reads_chinese_headings_paragraphs_and_tables(tmp_path, asynchronous):
    path = tmp_path / "学习资料.docx"
    document = WordDocument()
    document.add_heading("混合检索学习计划", level=1)
    document.add_paragraph("先学习向量检索，再比较全文检索。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "实验编号"
    table.cell(0, 1).text = "负责人"
    table.cell(1, 0).text = "HYB-2048"
    table.cell(1, 1).text = "周予安"
    document.add_paragraph("表格之后：每周二晚复习。")
    document.save(path)

    docs = await word_loader(str(path)) if asynchronous else word_loader_sync(str(path))

    assert len(docs) == 1
    text = docs[0].page_content
    assert "# 混合检索学习计划" in text
    assert "先学习向量检索，再比较全文检索。" in text
    assert "| 实验编号 | 负责人 |" in text
    assert "| HYB-2048 | 周予安 |" in text
    assert text.index("先学习") < text.index("实验编号") < text.index("表格之后")
    assert docs[0].metadata["source"] == str(path)
    assert "page" not in docs[0].metadata  # DOCX 没有稳定的排版页码


@pytest.mark.parametrize("asynchronous", [False, True])
async def test_word_loader_empty_document_returns_empty(tmp_path, asynchronous):
    path = tmp_path / "empty.docx"
    WordDocument().save(path)
    docs = await word_loader(str(path)) if asynchronous else word_loader_sync(str(path))
    assert docs == []


@pytest.mark.parametrize("asynchronous", [False, True])
async def test_word_loader_rejects_plain_text_disguised_as_docx(tmp_path, asynchronous):
    path = tmp_path / "corrupt.docx"
    path.write_text("这只是文本，不是有效的 Word 文件。", encoding="utf-8")
    docs = await word_loader(str(path)) if asynchronous else word_loader_sync(str(path))
    assert docs == []


# ---------------------------------------------------------------------------
# other loaders: error path only (no external parsers invoked)
# ---------------------------------------------------------------------------
def test_word_loader_sync_missing_file_returns_empty(tmp_path):
    assert word_loader_sync(str(tmp_path / "nope.docx")) == []


async def test_word_loader_async_missing_file_returns_empty(tmp_path):
    assert await word_loader(str(tmp_path / "nope.docx")) == []


def test_markdown_loader_sync_missing_file_returns_empty(tmp_path):
    assert markdown_loader_sync(str(tmp_path / "nope.md")) == []


async def test_markdown_loader_async_missing_file_returns_empty(tmp_path):
    assert await markdown_loader(str(tmp_path / "nope.md")) == []


def test_ppt_loader_sync_missing_file_returns_empty(tmp_path):
    assert ppt_loader_sync(str(tmp_path / "nope.pptx")) == []


async def test_ppt_loader_async_missing_file_returns_empty(tmp_path):
    assert await ppt_loader(str(tmp_path / "nope.pptx")) == []


# ---------------------------------------------------------------------------
# FontBBoxStreamFilter
# ---------------------------------------------------------------------------
def test_font_bbox_filter_suppresses_fontbox_lines():
    sink = io.StringIO()
    stream = FontBBoxStreamFilter(sink)
    stream.write("normal log line\n")
    stream.write("FontBBox from font descriptor (suppressed)\n")
    stream.write("more normal\n")
    out = sink.getvalue()
    assert "normal log line" in out
    assert "more normal" in out
    assert "FontBBox" not in out


if __name__ == "__main__":
    raise SystemExit(pytest.main([__file__, "-v"]))
